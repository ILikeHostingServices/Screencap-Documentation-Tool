#!/usr/bin/env python3
#
# test_pptx.py
# 2026-10-14
# Version: v1.0.0
#
# PURPOSE:
# Tests for the PowerPoint export (pptxwriter.py): the package is complete
# (every part has a content type, every relationship points at a part that
# exists, every picture reference resolves, all XML is well formed), the
# slides hold what they should (title slide, one slide per step, captions,
# alt text, speaker notes, clip slides, the logo), awkward input works
# (special characters, long and multi-line captions, JPEG screenshots, tall
# pictures, no steps, an SVG logo, the unblurred copy), and the command
# line and picture-size reading.
#
# python-pptx (pip install python-pptx) reads the files back like an Office
# program would. When SCREENCAP_OOXML_VALIDATOR is set (CI sets it), the
# files are also checked with Microsoft's Open XML SDK (tests/ooxml).
#
# Run from the repository root:
#   python -m unittest discover -s tests -v

import os
import re
import shlex
import shutil
import subprocess
import tempfile
import unittest
import zipfile
from pathlib import Path, PurePosixPath
from xml.etree import ElementTree as ET

from helpers import FFMPEG, FFPROBE, make_video  # noqa: E402

import export  # noqa: E402
import pptxwriter  # noqa: E402
import screencap as sc  # noqa: E402
import stepdoc  # noqa: E402

try:
    from pptx import Presentation
except ImportError:   # python-pptx is only needed by the tests
    Presentation = None

VALIDATOR = os.environ.get("SCREENCAP_OOXML_VALIDATOR")
REL = "{http://schemas.openxmlformats.org/package/2006/relationships}Relationship"
R_ATTR = re.compile(r'r:(?:embed|id)="([^"]+)"')


def check_package(test, path):
    """Fail the test when the .pptx is not a complete, consistent package."""
    with zipfile.ZipFile(path) as z:
        names = set(z.namelist())
        types = ET.fromstring(z.read("[Content_Types].xml"))
        defaults = {e.get("Extension").lower() for e in types
                    if e.tag.endswith("Default")}
        overrides = {e.get("PartName") for e in types if e.tag.endswith("Override")}
        for over in overrides:
            test.assertIn(over.lstrip("/"), names, f"content type for a missing part {over}")
        for name in names:
            if name.endswith("/"):
                continue
            ext = name.rsplit(".", 1)[-1].lower()
            test.assertTrue(f"/{name}" in overrides or ext in defaults,
                            f"{name} has no content type")
            if name.endswith((".xml", ".rels")):
                ET.fromstring(z.read(name))          # well formed
        for name in names:
            if not name.endswith(".rels"):
                continue
            folder = PurePosixPath(name).parent.parent
            source = folder / PurePosixPath(name).name[:-len(".rels")]
            ids = set()
            for rel in ET.fromstring(z.read(name)).iter(REL):
                ids.add(rel.get("Id"))
                target = os.path.normpath(str(folder / rel.get("Target"))).replace("\\", "/")
                test.assertIn(target, names, f"{name} points at missing {target}")
            if str(source) in names and str(source) != ".":
                used = set(R_ATTR.findall(z.read(str(source)).decode("utf-8")))
                test.assertLessEqual(used, ids, f"{source} uses undefined relationships")
        for name in names:
            if name.startswith("ppt/slides/slide") and name.endswith(".xml"):
                shape_ids = re.findall(r'<p:cNvPr id="(\d+)"', z.read(name).decode("utf-8"))
                test.assertEqual(len(shape_ids), len(set(shape_ids)), f"{name}: shape ids repeat")


def validate_with_sdk(test, *paths):
    if not VALIDATOR:
        return
    res = subprocess.run(shlex.split(VALIDATOR) + [str(p) for p in paths],
                         stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True)
    test.assertEqual(res.returncode, 0, res.stdout)


class PictureSizeTests(unittest.TestCase):
    @unittest.skipUnless(FFMPEG, "FFmpeg is not installed")
    def test_png_gif_and_jpeg_sizes(self):
        with tempfile.TemporaryDirectory() as tmp:
            for name, extra in (("a.png", []), ("b.gif", []), ("c.jpg", []),
                                ("d.jpg", ["-q:v", "2", "-pix_fmt", "yuvj444p"])):
                path = Path(tmp) / name
                subprocess.run([FFMPEG, "-loglevel", "error", "-y", "-f", "lavfi", "-i",
                                "color=red:s=322x124", "-frames:v", "1"] + extra + [str(path)],
                               check=True)
                self.assertEqual(pptxwriter.image_size(path), (322, 124), name)
            bad = Path(tmp) / "x.svg"
            bad.write_text("<svg/>")
            with self.assertRaises(ValueError):
                pptxwriter.image_size(bad)

    def test_fit_keeps_the_shape(self):
        self.assertEqual(pptxwriter.fit(1920, 1080, 960, 1000), (960, 540))
        self.assertEqual(pptxwriter.fit(400, 2000, 1000, 500), (100, 500))


@unittest.skipUnless(FFMPEG and FFPROBE, "FFmpeg is not installed")
class PowerPointExportTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = Path(tempfile.mkdtemp(prefix="screencap_pptx_"))
        source = cls.tmp / "source"
        source.mkdir()
        make_video(source / "Demo.mp4")
        assert sc.main(["-s", str(source), "-o", str(cls.tmp / "out"), "--no-redact"]) == 0
        cls.folder = cls.tmp / "out" / "Demo"
        cls.logo = cls.tmp / "logo.png"
        subprocess.run([FFMPEG, "-loglevel", "error", "-y", "-f", "lavfi", "-i",
                        "color=0x2563eb:s=300x90", "-frames:v", "1", str(cls.logo)], check=True)

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.tmp, ignore_errors=True)

    def doc(self):
        doc = stepdoc.load(self.folder)
        doc["document"].update({"title": 'Set up <VPN> & "Mail" – café \U0001F512',
                                "version": "v2.0.0", "author": "IT Service Desk"})
        return doc

    def export(self, doc, **kw):
        dest = export.export(self.folder, doc, "pptx", **kw)
        self.assertEqual(dest.suffix, ".pptx")
        check_package(self, dest)
        return dest

    @unittest.skipUnless(Presentation, "python-pptx is not installed")
    def test_slides_hold_the_steps(self):
        doc = self.doc()
        steps = doc["steps"]
        steps[0]["caption"] = "Open the installer.\nThen click Next.\x07"
        steps[0]["narration"] = "First, open the installer from the share."
        steps[1]["caption"] = "Long caption. " * 40
        dest = self.export(doc, logo=self.logo)
        validate_with_sdk(self, dest)
        prs = Presentation(str(dest))
        self.assertEqual((prs.slide_width, prs.slide_height), (12192000, 6858000))
        slides = list(prs.slides)
        self.assertEqual(len(slides), len(steps) + 1)
        title = " ".join(s.text_frame.text for s in slides[0].shapes if s.has_text_frame)
        self.assertIn('Set up <VPN> & "Mail" – café \U0001F512', title)
        self.assertIn("Version v2.0.0", title)
        self.assertIn("IT Service Desk", title)
        for i, (slide, step) in enumerate(zip(slides[1:], steps), 1):
            texts = [s.text_frame.text for s in slide.shapes if s.has_text_frame]
            self.assertTrue(texts[0].startswith(f"Step {i}"))
            pics = [s for s in slide.shapes if s.shape_type == 13]
            self.assertEqual([p.name for p in pics], ["Logo", "Screenshot"])
            shot = pics[1]
            descr = shot._element.xpath("./p:nvPicPr/p:cNvPr/@descr")[0]
            self.assertEqual(descr, stepdoc.alt_text(step, i).replace("\x07", ""))
            self.assertLessEqual(shot.left + shot.width, prs.slide_width)
            self.assertLessEqual(shot.top + shot.height, prs.slide_height)
            if step.get("caption"):
                caption = texts[-1]
                self.assertTrue(caption.startswith(step["caption"].strip()[:20]))
                cap = [s for s in slide.shapes if s.name == "Caption"][0]
                self.assertLessEqual(cap.top + cap.height, prs.slide_height)
                self.assertGreaterEqual(cap.top, shot.top + shot.height)
        self.assertEqual(slides[1].notes_slide.notes_text_frame.text,
                         "First, open the installer from the share.")
        self.assertFalse(slides[2].has_notes_slide)
        self.assertEqual(slides[1].shapes[-1].text_frame.text,
                         "Open the installer.\nThen click Next.")
        with zipfile.ZipFile(dest) as z:
            media = [n for n in z.namelist() if n.startswith("ppt/media/")]
        self.assertEqual(len(media), len(steps) + 1, "the logo is stored once")

    @unittest.skipUnless(Presentation, "python-pptx is not installed")
    def test_clip_slide_jpeg_and_tall_pictures(self):
        doc = self.doc()
        step = doc["steps"][1]
        step["clip"] = True
        stepdoc.render(self.folder, doc, FFMPEG)
        tall = self.tmp / "tall.jpg"
        subprocess.run([FFMPEG, "-loglevel", "error", "-y", "-f", "lavfi", "-i",
                        "color=gray:s=300x1600", "-frames:v", "1", str(tall)], check=True)
        deck = pptxwriter.Deck()
        pptxwriter.step_slide(deck, 1, "00:00:01", tall, "Tall", "A caption")
        dest = deck.write(self.tmp / "tall.pptx")
        check_package(self, dest)
        validate_with_sdk(self, dest)
        prs = Presentation(str(dest))
        pic = [s for s in prs.slides[0].shapes if s.shape_type == 13][0]
        self.assertLessEqual(pic.top + pic.height, prs.slide_height)
        self.assertAlmostEqual(pic.width / pic.height, 300 / 1600, places=2)

        dest = self.export(doc)
        validate_with_sdk(self, dest)
        slides = list(Presentation(str(dest)).slides)
        self.assertEqual(len(slides), len(doc["steps"]) + 2)
        clip_slide = slides[3]
        self.assertEqual(clip_slide.shapes[0].text_frame.text, "Step 2: animation")
        with zipfile.ZipFile(dest) as z:
            self.assertTrue(any(n.endswith(".gif") for n in z.namelist()))
            self.assertIn('Extension="gif" ContentType="image/gif"',
                          z.read("[Content_Types].xml").decode())
        step["clip"] = False
        stepdoc.render(self.folder, doc, FFMPEG)

    @unittest.skipUnless(Presentation, "python-pptx is not installed")
    def test_no_steps_svg_logo_and_unblurred_copy(self):
        doc = self.doc()
        doc["steps"] = []
        svg = self.tmp / "logo.svg"
        svg.write_text('<svg xmlns="http://www.w3.org/2000/svg" width="10" height="10"/>')
        with self.assertLogs("screencap", level="WARNING") as logs:
            dest = self.export(doc, logo=svg)
        self.assertIn("cannot show the SVG logo", "\n".join(logs.output))
        self.assertEqual(len(Presentation(str(dest)).slides), 1)
        validate_with_sdk(self, dest)

        with self.assertRaisesRegex(RuntimeError, "no unblurred screenshots"):
            export.export(self.folder, self.doc(), "pptx", unredacted=True)
        (self.folder / stepdoc.UNREDACTED_DIR).mkdir(exist_ok=True)
        doc = self.doc()
        for step in doc["steps"]:
            shutil.copy(self.folder / step["file"],
                        self.folder / stepdoc.UNREDACTED_DIR / step["file"])
        dest = export.export(self.folder, doc, "pptx", unredacted=True)
        check_package(self, dest)
        self.assertIn("UNREDACTED", dest.name)
        first = Presentation(str(dest)).slides[0]
        self.assertIn("UNREDACTED COPY",
                      " ".join(s.text_frame.text for s in first.shapes if s.has_text_frame))
        shutil.rmtree(self.folder / stepdoc.UNREDACTED_DIR)

    def test_command_line_export(self):
        out = self.tmp / "cli"
        source = self.tmp / "source"
        self.assertEqual(sc.main(["-s", str(source), "-o", str(out), "--no-redact",
                                  "--export", "html,pptx"]), 0)
        files = sorted(p.suffix for p in (out / "Demo" / "export").iterdir())
        self.assertEqual(files, [".html", ".pptx"])
        check_package(self, next((out / "Demo" / "export").glob("*.pptx")))


if __name__ == "__main__":
    unittest.main()
